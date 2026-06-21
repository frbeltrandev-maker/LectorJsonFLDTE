import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


NOMBRE_ARCHIVO_INVALIDOS = "InvalidosDTE.xlsx"
NOMBRE_HOJA_INVALIDOS = "DTE Invalidos"
ENCABEZADOS_INVALIDOS = [
    "Nombre del Archivo",
    "Codigo Generacion",
    "Fecha Emision",
    "Estado DTE",
    "Descripcion de Hacienda",
]


def guardar_dtes_invalidos(salida, dtes_invalidos):
    if not dtes_invalidos:
        return

    ruta = os.path.join(salida, NOMBRE_ARCHIVO_INVALIDOS)
    if os.path.exists(ruta):
        wb = load_workbook(ruta)
        ws = wb[NOMBRE_HOJA_INVALIDOS] if NOMBRE_HOJA_INVALIDOS in wb.sheetnames else wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = NOMBRE_HOJA_INVALIDOS
        ws.append(ENCABEZADOS_INVALIDOS)
        for cell in ws[1]:
            cell.font = Font(bold=True)

    for dte_invalido in dtes_invalidos:
        ws.append(dte_invalido)

    wb.save(ruta)
