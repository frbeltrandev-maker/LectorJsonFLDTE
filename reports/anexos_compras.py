import os

from openpyxl import Workbook
from openpyxl.styles import Font

from dte.calculations import calcular_valores_resumen, iterar_tributos_conocidos
from dte.extractor import (
    obtener_contenido_dte,
    obtener_fecha_formateada_dte,
    obtener_sello_recibido,
)
from dte.reader import cargar_datos_archivo
from dte.types import (
    DTE_COMPRA_CON_FACTURA,
    DTE_COMPROBANTE_CREDITO_FISCAL,
    DTE_FACTURA,
    DTE_NOTA_CREDITO,
    DTE_NOTA_DEBITO,
)
from services.hacienda import (
    ConsultorHaciendaPorLotes,
    ConsultaHaciendaError,
    LimiteConsultasHaciendaError,
    consultar_estado_dte,
    dte_esta_transmitido_satisfactoriamente,
)
from reports.invalidos_dte import guardar_dtes_invalidos


def _crear_contadores():
    return {
        "sum_factura": 0,
        "sum_ntcdebito": 0,
        "sum_ccf": 0,
        "sum_ntccredito": 0,
        "sum_grav": 0.0,
        "sum_iva": 0.0,
        "sum_total": 0.0,
        "sum_fovial": 0.0,
        "sum_cotrans": 0.0,
        "sum_docjson": 0,
    }


def _crear_valores():
    return {
        "fovial": 0,
        "cotrans": 0,
        "total_exenta": 0,
        "total_gravada": 0,
        "total_no_sujeta": 0,
        "iva": 0,
        "monto_total_operacion": 0,
        "iva_perci1": 0,
    }


def generar_anexo_compras(
    carpeta,
    salida,
    nombre_archivo_salida,
    consultar_estado=consultar_estado_dte,
    progreso=None,
    estado=None,
):
    contadores = _crear_contadores()
    archivos_no_json = []
    dtes_invalidos = []
    consultas_hacienda = {}
    def avisar_espera(segundos, motivo):
        if estado:
            estado(
                f"{motivo} Pausa automática de {segundos} segundos para "
                "respetar el límite de Hacienda."
            )

    def avisar_reanudacion():
        if estado:
            estado("Reanudando las consultas con Hacienda...")

    def consultar_con_estado(codigo_generacion, fecha_emi):
        if consultar_estado is consultar_estado_dte:
            return consultar_estado(
                codigo_generacion,
                fecha_emi,
                al_esperar=avisar_espera,
                al_reanudar=avisar_reanudacion,
            )
        return consultar_estado(codigo_generacion, fecha_emi)

    consultor_hacienda = ConsultorHaciendaPorLotes(
        consultar_con_estado,
        al_esperar=avisar_espera,
        al_reanudar=avisar_reanudacion,
    )
    consultar_hacienda = True
    motivo_sin_verificacion = None
    wb = Workbook()
    ws_resumen = wb.active
    ws_resumen.title = "Resumen de DTEs"
    ws_percepcion = wb.create_sheet(title="Percepción")
    ws_errores = wb.create_sheet(title="Errores")

    ws_resumen.append([
        "Fecha", "Clase de Documento", "Tipo de DTE", "Num Doc", "NRC", "Nombre de Proveedor",
        "Exenta/NoSuj", "IntE/S", "ImpE/S", "Internas Gravadas", "Internaciones Gravadas",
        "Importaciones Gravadas de Bienes", "Importaciones Gravadas de Servicio",
        "Credito Fiscal", "Monto Total", "Fovial", "Cotrans", "Percibido 1%", "Num_Anexo",
    ])
    ws_percepcion.append([
        "NIT", "Fecha", "Tipo DTE", "Sello Recepcion", "CodGeneracion", "Monto Sujeto",
        "Monto de Percepción", "DUI", "Num_Anexo",
    ])
    ws_errores.append(["Nombre del Archivo", "Error o aviso"])

    for hoja in (ws_resumen, ws_percepcion, ws_errores):
        for cell in hoja[1]:
            cell.font = Font(bold=True)

    nombres_archivos = os.listdir(carpeta)
    total_archivos = len(nombres_archivos)

    for indice, nombre_archivo in enumerate(nombres_archivos, start=1):
        if progreso:
            progreso(indice, total_archivos, nombre_archivo)

        ruta_archivo = os.path.join(carpeta, nombre_archivo)
        if not os.path.isfile(ruta_archivo) or not nombre_archivo.lower().endswith(".json"):
            archivos_no_json.append((nombre_archivo, "No es un archivo JSON"))
            continue

        print(f"Procesando archivo: {nombre_archivo}")
        datos = cargar_datos_archivo(ruta_archivo)
        if not datos:
            continue

        contenido = obtener_contenido_dte(datos, "emisor")
        if not contenido:
            print(f"Estructura no valida en el archivo {nombre_archivo}")
            continue

        valores = _crear_valores()
        emisor = contenido["emisor"]
        identificacion = contenido["identificacion"]
        resumen = contenido["resumen"]
        tipo_dte = identificacion.get("tipoDte")
        fecha_emi = identificacion.get("fecEmi")
        _, fecha_formateada = obtener_fecha_formateada_dte(fecha_emi)

        if tipo_dte in DTE_COMPRA_CON_FACTURA:
            codigo_generacion = identificacion.get("codigoGeneracion")
            clave_consulta = (codigo_generacion, fecha_emi)
            respuesta_hacienda = consultas_hacienda.get(clave_consulta)
            if respuesta_hacienda is None and consultar_hacienda:
                try:
                    respuesta_hacienda = consultor_hacienda.consultar(
                        codigo_generacion,
                        fecha_emi,
                    )
                    consultas_hacienda[clave_consulta] = respuesta_hacienda
                except LimiteConsultasHaciendaError as error:
                    consultar_hacienda = False
                    motivo_sin_verificacion = str(error)
                except ConsultaHaciendaError as error:
                    archivos_no_json.append((
                        nombre_archivo,
                        f"No se pudo consultar Hacienda: {error}",
                    ))
                    continue

            if respuesta_hacienda is None:
                archivos_no_json.append((
                    nombre_archivo,
                    "Procesado sin verificación en Hacienda: "
                    f"{motivo_sin_verificacion}",
                ))
            elif not dte_esta_transmitido_satisfactoriamente(respuesta_hacienda):
                estado_documento = respuesta_hacienda.get("estadoDoc", "Sin estado")
                dtes_invalidos.append([
                    nombre_archivo,
                    codigo_generacion,
                    fecha_formateada,
                    estado_documento,
                    respuesta_hacienda.get("descripcionEstado", ""),
                ])
                continue

            valores.update(calcular_valores_resumen(
                resumen,
                campo_no_sujeta="totalNoSuj",
                campo_monto_total="montoTotalOperacion",
            ))
            for clave, valor in iterar_tributos_conocidos(resumen):
                valores[clave] = valor
                contadores[f"sum_{clave}"] += valor

            if tipo_dte == DTE_COMPROBANTE_CREDITO_FISCAL:
                contadores["sum_ccf"] += 1
                contadores["sum_grav"] += valores["total_gravada"]
                contadores["sum_total"] += valores["monto_total_operacion"]
            if tipo_dte == DTE_NOTA_CREDITO:
                contadores["sum_ntccredito"] += 1
                contadores["sum_grav"] -= valores["total_gravada"]
                contadores["sum_total"] -= valores["monto_total_operacion"]
            if tipo_dte == DTE_NOTA_DEBITO:
                contadores["sum_ntcdebito"] += 1
                contadores["sum_grav"] += valores["total_gravada"]
                contadores["sum_total"] += valores["monto_total_operacion"]
            if tipo_dte == DTE_FACTURA:
                contadores["sum_factura"] += 1

            ws_resumen.append([
                fecha_formateada,
                "4",
                tipo_dte,
                codigo_generacion,
                emisor.get("nrc"),
                emisor.get("nombre"),
                valores["total_exenta"] + valores["total_no_sujeta"],
                0,
                0,
                valores["total_gravada"],
                0,
                0,
                0,
                valores["iva"],
                valores["monto_total_operacion"],
                valores["fovial"],
                valores["cotrans"],
                valores["iva_perci1"],
                3,
            ])

        contadores["sum_docjson"] += 1

        if valores["iva_perci1"] > 0:
            ws_percepcion.append([
                emisor.get("nit"),
                fecha_formateada,
                tipo_dte,
                obtener_sello_recibido(contenido),
                codigo_generacion,
                valores["total_gravada"],
                valores["iva_perci1"],
                "",
                8,
            ])

    wb.save(os.path.join(salida, f"{nombre_archivo_salida}.xlsx"))

    if archivos_no_json:
        for archivo in archivos_no_json:
            ws_errores.append(archivo)
        wb.save(os.path.join(salida, f"{nombre_archivo_salida}_errores.xlsx"))

    guardar_dtes_invalidos(salida, dtes_invalidos)

    return [
        ("Facturas", contadores["sum_factura"]),
        ("Notas de Crédito", contadores["sum_ntccredito"]),
        ("Notas de Débito", contadores["sum_ntcdebito"]),
        ("Comprobantes de Crédito Fiscal", contadores["sum_ccf"]),
        ("Documentos Procesados", contadores["sum_docjson"]),
    ]
