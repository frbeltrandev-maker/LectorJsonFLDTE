"""Anexo 6 de anticipos del 2%, según la estructura suministrada para DTE-09."""

import os
import re
from decimal import Decimal, InvalidOperation, ROUND_DOWN

from openpyxl import Workbook

from dte.duplicates import ControlDuplicadosDTE
from dte.extractor import obtener_contenido_dte, obtener_fecha_formateada_dte, obtener_sello_recibido
from dte.reader import cargar_datos_archivo
from dte.types import DTE_DOCUMENTO_CONTABLE_LIQUIDACION
from services.hacienda import (
    ConsultorHaciendaPorLotes,
    ConsultaHaciendaError,
    LimiteConsultasHaciendaError,
    consultar_estado_dte,
    dte_esta_transmitido_satisfactoriamente,
)


def _importe(cuerpo, campo):
    try:
        valor = Decimal(str(cuerpo.get(campo)))
        if not valor.is_finite() or valor < 0:
            raise ValueError(f"{campo}: el importe debe ser finito y no negativo")
        # La referencia indica que se toman únicamente dos decimales.
        valor = valor.quantize(Decimal("0.01"), rounding=ROUND_DOWN)
    except (InvalidOperation, TypeError):
        raise ValueError(f"{campo}: importe ausente o inválido") from None
    if len(format(valor, ".2f")) > 10:
        raise ValueError(f"{campo}: excede los 10 caracteres permitidos")
    return valor


def _fila_documento(contenido):
    identificacion = contenido["identificacion"]
    emisor = contenido["emisor"]
    cuerpo = contenido["cuerpoDocumento"]
    if not all(isinstance(valor, dict) for valor in (identificacion, emisor, cuerpo)):
        raise ValueError("Identificación, emisor o cuerpoDocumento inválidos")
    fecha, fecha_texto = obtener_fecha_formateada_dte(identificacion.get("fecEmi"))
    if fecha is None:
        raise ValueError("Falta la fecha de emisión")
    documento_agente = re.sub(r"[-/\s]", "", str(emisor.get("nit") or ""))
    nit, dui = "", ""
    if re.fullmatch(r"[0-9]{14}", documento_agente):
        nit = documento_agente
    elif re.fullmatch(r"[0-9]{9}", documento_agente) and fecha.year >= 2022:
        dui = documento_agente
    else:
        raise ValueError("El agente necesita NIT de 14 dígitos o DUI de 9 desde enero de 2022")
    codigo = identificacion.get("codigoGeneracion")
    if not isinstance(codigo, str) or not codigo.strip():
        raise ValueError("Falta codigoGeneracion")
    # Excepción histórica indicada en la referencia para noviembre de 2022.
    numero = identificacion.get("numeroControl") if (fecha.year, fecha.month) == (2022, 11) else codigo
    numero = str(numero or "").strip().replace("-", "").upper()
    if not numero or len(numero) > 100:
        raise ValueError("Número de documento ausente o mayor de 100 caracteres")
    monto = _importe(cuerpo, "montoSujetoPercepcion")
    anticipo = _importe(cuerpo, "ivaPercibido")
    return [nit, fecha_texto, "", numero, monto, anticipo, dui, 6]


def generar_anexo_anticipo_iva(
    carpeta, salida, nombre_archivo_salida,
    consultar_estado=consultar_estado_dte, progreso=None, estado=None,
):
    duplicados = ControlDuplicadosDTE()
    errores = []
    cantidad = 0
    motivo_sin_verificacion = None

    def avisar_espera(segundos, motivo):
        if estado:
            estado(f"{motivo} Pausa automática de {segundos} segundos.")

    def avisar_reanudacion():
        if estado:
            estado("Reanudando las consultas con Hacienda...")

    def consultar(codigo, fecha):
        if consultar_estado is consultar_estado_dte:
            return consultar_estado(codigo, fecha, al_esperar=avisar_espera, al_reanudar=avisar_reanudacion)
        return consultar_estado(codigo, fecha)

    consultor = ConsultorHaciendaPorLotes(
        consultar, al_esperar=avisar_espera, al_reanudar=avisar_reanudacion,
    )
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Anticipo IVA 2%"
    archivos = sorted(os.listdir(carpeta), key=str.casefold)
    for indice, nombre in enumerate(archivos, start=1):
        if progreso:
            progreso(indice, len(archivos), nombre)
        ruta = os.path.join(carpeta, nombre)
        if not os.path.isfile(ruta) or not nombre.lower().endswith(".json"):
            continue
        datos = cargar_datos_archivo(ruta)
        contenido = obtener_contenido_dte(datos, "emisor", "cuerpoDocumento")
        if not contenido:
            # Otras clases de DTE no pertenecen a este anexo.
            candidatos = [datos] + list(datos.values()) if isinstance(datos, dict) else []
            if any(isinstance(d, dict) and isinstance(d.get("identificacion"), dict)
                   and d["identificacion"].get("tipoDte") not in (None, "09") for d in candidatos):
                continue
            errores.append((nombre, "JSON inválido o faltan campos del DTE-09"))
            continue
        identificacion = contenido.get("identificacion")
        if not isinstance(identificacion, dict):
            errores.append((nombre, "Identificación inválida"))
            continue
        if identificacion.get("tipoDte") != DTE_DOCUMENTO_CONTABLE_LIQUIDACION:
            continue
        try:
            fila = _fila_documento(contenido)
        except (ValueError, TypeError) as error:
            errores.append((nombre, str(error)))
            continue
        if duplicados.es_duplicado(identificacion, nombre):
            continue
        respuesta = None
        if motivo_sin_verificacion is None:
            try:
                respuesta = consultor.consultar(identificacion["codigoGeneracion"], identificacion["fecEmi"])
            except LimiteConsultasHaciendaError as error:
                motivo_sin_verificacion = str(error)
            except ConsultaHaciendaError as error:
                errores.append((nombre, f"No se pudo consultar Hacienda: {error}"))
                continue
        if respuesta is not None and not dte_esta_transmitido_satisfactoriamente(respuesta):
            errores.append((nombre, f"DTE excluido por estado en Hacienda: {respuesta.get('estadoDoc', 'Sin estado')}"))
            continue
        sello = ((respuesta or {}).get("selloVal") or obtener_sello_recibido(contenido)
                 or obtener_sello_recibido(datos))
        if not isinstance(sello, str) or not sello.strip() or len(sello.strip()) > 100:
            errores.append((nombre, "Sello de recepción ausente o inválido"))
            continue
        if motivo_sin_verificacion is not None:
            errores.append((nombre, f"Procesado sin verificación en Hacienda: {motivo_sin_verificacion}"))
        fila[2] = sello.strip()
        hoja.append(fila)
        for columna in (1, 2, 3, 4, 7):
            celda = hoja.cell(hoja.max_row, columna)
            celda.number_format = "@"
            celda.data_type = "s"
        for columna in (5, 6):
            hoja.cell(hoja.max_row, columna).number_format = "0.00"
        cantidad += 1

    libro.save(os.path.join(salida, f"{nombre_archivo_salida}.xlsx"))
    errores.extend(duplicados.avisos)
    if errores:
        libro_errores = Workbook()
        libro_errores.active.title = "Errores"
        libro_errores.active.append(["Nombre del Archivo", "Error o aviso"])
        for error in errores:
            libro_errores.active.append(error)
        libro_errores.save(os.path.join(salida, f"{nombre_archivo_salida}_errores.xlsx"))
    return [("Documentos Contables de Liquidación", cantidad), ("Documentos Procesados", cantidad)]
