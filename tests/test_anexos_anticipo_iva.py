import copy
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock

from openpyxl import load_workbook

from dte.types import TIPO_DOCUMENTO_ANTICIPO_IVA, TIPOS_DOCUMENTO
from reports.anexos_anticipo_iva import generar_anexo_anticipo_iva
from services.hacienda import LimiteConsultasHaciendaError


def documento():
    return {
        "identificacion": {
            "tipoDte": "09", "fecEmi": "2026-10-01",
            "codigoGeneracion": "ABCDEF01-1234-4321-ABCD-123456789012",
            "numeroControl": "DTE-09-M001P001-000000000000001",
        },
        "emisor": {"nit": "0614-123456-123-4"},
        "receptor": {"nit": "99999999999999"},
        "cuerpoDocumento": {
            "valorOperaciones": 2000, "subTotal": 1695, "liquidoApagar": 1600,
            "montoSujetoPercepcion": 1500, "ivaPercibido": 30,
        },
        "selloRecibido": "A" * 40,
    }


class AnticipoIVATests(unittest.TestCase):
    def setUp(self):
        temporal = tempfile.TemporaryDirectory()
        self.addCleanup(temporal.cleanup)
        self.raiz = Path(temporal.name)
        self.entrada = self.raiz / "entrada"
        self.entrada.mkdir()
        self.consulta = Mock(return_value={
            "estadoDoc": "Transmitido Satisfactoriamente", "selloVal": "B" * 40,
        })

    def escribir(self, nombre, datos):
        (self.entrada / nombre).write_text(json.dumps(datos), encoding="utf-8")

    def generar(self):
        resultado = generar_anexo_anticipo_iva(
            self.entrada, self.raiz, "anticipo", consultar_estado=self.consulta,
        )
        libro = load_workbook(self.raiz / "anticipo.xlsx")
        self.addCleanup(libro.close)
        return resultado, libro

    def test_estructura_sin_encabezados_emisor_importes_y_sello_hacienda(self):
        self.escribir("dte.json", documento())
        resultado, libro = self.generar()
        self.assertEqual(resultado[-1], ("Documentos Procesados", 1))
        self.assertEqual(libro.sheetnames, ["Anticipo IVA 2%"])
        hoja = libro.active
        self.assertEqual((hoja.max_row, hoja.max_column), (1, 8))
        self.assertEqual(list(hoja.values)[0], (
            "06141234561234", "01/10/2026", "B" * 40,
            "ABCDEF0112344321ABCD123456789012", 1500, 30, None, 6,
        ))
        self.assertEqual(len(hoja.merged_cells.ranges), 0)
        self.assertEqual(hoja["A1"].data_type, "s")
        self.assertEqual(hoja["E1"].number_format, "0.00")
        self.assertEqual(hoja["F1"].number_format, "0.00")
        self.consulta.assert_called_once_with(documento()["identificacion"]["codigoGeneracion"], "2026-10-01")

    def test_anidado_dui_sello_externo_duplicados_y_otros_tipos(self):
        dte = documento()
        dte["emisor"]["nit"] = "01234567-8"
        del dte["selloRecibido"]
        datos = {"dte": dte, "respuestaHacienda": {"selloRecibido": "C" * 40}}
        self.escribir("01.json", datos)
        copia = copy.deepcopy(datos)
        copia["dte"]["identificacion"]["codigoGeneracion"] = dte["identificacion"]["codigoGeneracion"].lower()
        self.escribir("02-copia.json", copia)
        self.escribir("03-otro.json", {"identificacion": {"tipoDte": "03"}, "emisor": {}, "resumen": {}})
        self.consulta.return_value = {"estadoDoc": "Transmitido Satisfactoriamente"}
        resultado, libro = self.generar()
        self.assertEqual(resultado[-1][1], 1)
        self.assertIsNone(libro.active["A1"].value)
        self.assertEqual(libro.active["G1"].value, "012345678")
        self.assertEqual(libro.active["C1"].value, "C" * 40)
        self.assertEqual(self.consulta.call_count, 1)
        errores = load_workbook(self.raiz / "anticipo_errores.xlsx")
        try:
            self.assertEqual(errores.active.max_row, 2)
            self.assertIn("duplicado", errores.active["B2"].value)
        finally:
            errores.close()

    def test_429_conserva_importes_y_sello_local(self):
        self.escribir("dte.json", documento())
        self.consulta.side_effect = LimiteConsultasHaciendaError("HTTP 429")
        resultado, libro = self.generar()
        self.assertEqual(resultado[-1][1], 1)
        self.assertEqual(libro.active["C1"].value, "A" * 40)
        self.assertEqual(libro.active["F1"].value, 30)
        self.assertTrue((self.raiz / "anticipo_errores.xlsx").exists())

    def test_excluye_documento_invalidado(self):
        self.escribir("dte.json", documento())
        self.consulta.return_value = {"estadoDoc": "Invalidado", "selloVal": "B" * 40}
        resultado, libro = self.generar()
        self.assertEqual(resultado[-1][1], 0)
        self.assertIsNone(libro.active["A1"].value)
        self.assertTrue((self.raiz / "anticipo_errores.xlsx").exists())

    def test_campos_invalidos_se_reportan_sin_detener_lote(self):
        for indice, (campo, valor) in enumerate((
            ("montoSujetoPercepcion", -1), ("ivaPercibido", None),
            ("ivaPercibido", 100000000), ("ivaPercibido", "NaN"),
        )):
            dte = documento()
            dte["cuerpoDocumento"][campo] = valor
            self.escribir(f"{indice}.json", dte)
        self.escribir("valido.json", documento())
        resultado, libro = self.generar()
        self.assertEqual(resultado[-1][1], 1)
        self.assertEqual(libro.active["E1"].value, 1500)
        self.assertEqual(self.consulta.call_count, 1)
        errores = load_workbook(self.raiz / "anticipo_errores.xlsx")
        try:
            self.assertEqual(errores.active.max_row, 5)
        finally:
            errores.close()

    def test_noviembre_2022_numero_control_y_dos_decimales(self):
        dte = documento()
        dte["identificacion"]["fecEmi"] = "2022-11-01"
        dte["cuerpoDocumento"]["montoSujetoPercepcion"] = 1500.129
        dte["cuerpoDocumento"]["ivaPercibido"] = 30.129
        self.escribir("dte.json", dte)
        _, libro = self.generar()
        self.assertEqual(libro.active["D1"].value, "DTE09M001P001000000000000001")
        self.assertAlmostEqual(libro.active["E1"].value, 1500.12)
        self.assertAlmostEqual(libro.active["F1"].value, 30.12)

    def test_opcion_disponible_en_interfaz(self):
        from lector_carpetajson import REPORTES
        self.assertIn(TIPO_DOCUMENTO_ANTICIPO_IVA, TIPOS_DOCUMENTO)
        self.assertIs(REPORTES[TIPO_DOCUMENTO_ANTICIPO_IVA][0], generar_anexo_anticipo_iva)
