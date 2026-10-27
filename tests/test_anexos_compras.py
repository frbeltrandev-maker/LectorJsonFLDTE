import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from reports.anexos_compras import generar_anexo_compras
from services.hacienda import LimiteConsultasHaciendaError
from tests.dte_fixtures import crear_dte
from unittest.mock import Mock


class AnexosComprasTests(unittest.TestCase):
    def test_percepcion_prioriza_sello_hacienda_y_recurre_al_local(self):
        for sello_mh in ("SELLO-HACIENDA", None, ""):
            for anidado in (False, True):
                with self.subTest(sello=sello_mh, anidado=anidado), tempfile.TemporaryDirectory() as directorio:
                    carpeta = Path(directorio)
                    dte = crear_dte()
                    dte["resumen"]["ivaPerci1"] = 0.03
                    datos = {"dte": dte} if anidado else dte
                    (carpeta / "dte.json").write_text(json.dumps(datos), encoding="utf-8")
                    consulta = Mock(return_value={
                        "estadoDoc": "Transmitido Satisfactoriamente",
                        "selloVal": sello_mh,
                    })
                    generar_anexo_compras(carpeta, carpeta, "compras", consultar_estado=consulta)
                    consulta.assert_called_once_with("GEN-001", "2026-10-01")
                    libro = load_workbook(carpeta / "compras.xlsx")
                    try:
                        fila = list(libro["Percepción"].values)[1]
                        self.assertEqual(fila[3], sello_mh or "SELLO-PRUEBA")
                        self.assertEqual(fila[4], "GEN-001")
                        self.assertEqual(fila[5], 3)
                        self.assertAlmostEqual(fila[6], 0.03)
                    finally:
                        libro.close()

    def test_percepcion_usa_sello_local_tras_limite_sin_reutilizar_sello_anterior(self):
        with tempfile.TemporaryDirectory() as directorio:
            carpeta = Path(directorio)
            for indice in range(1, 4):
                dte = crear_dte(codigo=f"GEN-{indice}")
                dte["resumen"]["ivaPerci1"] = 0.03
                dte["selloRecibido"] = f"SELLO-LOCAL-{indice}"
                (carpeta / f"{indice}.json").write_text(json.dumps(dte), encoding="utf-8")
            consulta = Mock(side_effect=[
                {"estadoDoc": "Transmitido Satisfactoriamente", "selloVal": "SELLO-MH-1"},
                LimiteConsultasHaciendaError("HTTP 429"),
            ])
            resumen = dict(generar_anexo_compras(carpeta, carpeta, "compras", consultar_estado=consulta))
            self.assertEqual(consulta.call_count, 2)
            self.assertEqual(resumen["Documentos Procesados"], 3)
            libro = load_workbook(carpeta / "compras.xlsx")
            try:
                filas = list(libro["Percepción"].values)[1:]
                self.assertEqual([fila[3] for fila in filas], ["SELLO-MH-1", "SELLO-LOCAL-2", "SELLO-LOCAL-3"])
                self.assertAlmostEqual(sum(fila[6] for fila in filas), 0.09)
                self.assertEqual(libro["Errores"].max_row, 3)
            finally:
                libro.close()

    def test_genera_anexo_y_resumen_de_compras(self):
        with tempfile.TemporaryDirectory() as directorio:
            carpeta = Path(directorio)
            dte = {
                "identificacion": {
                    "fecEmi": "2026-06-21",
                    "tipoDte": "03",
                    "codigoGeneracion": "GEN-001",
                },
                "emisor": {
                    "nit": "0614-000000-000-0",
                    "nrc": "12345",
                    "nombre": "Proveedor de Prueba",
                },
                "resumen": {
                    "totalExenta": 10.00,
                    "descuExenta": 2.00,
                    "totalNoSuj": 4.00,
                    "descuNoSuj": 1.00,
                    "totalGravada": 20.00,
                    "descuGravada": 5.00,
                    "montoTotalOperacion": 30.00,
                    "ivaPerci1": 0.30,
                    "tributos": [
                        {"codigo": "D1", "valor": 1.00},
                        {"codigo": "C8", "valor": 2.00},
                        {"codigo": "20", "valor": 3.00},
                    ],
                },
                "selloRecibido": "SELLO-001",
            }
            (carpeta / "dte.json").write_text(json.dumps(dte), encoding="utf-8")
            (carpeta / "nota.txt").write_text("archivo auxiliar", encoding="utf-8")

            progresos = []
            resumen = generar_anexo_compras(
                carpeta,
                carpeta,
                "compras",
                consultar_estado=lambda codigo, fecha: {
                    "estadoDoc": "Transmitido Satisfactoriamente",
                },
                progreso=lambda actual, total, archivo: progresos.append((actual, total, archivo)),
            )

            self.assertEqual(
                resumen,
                [
                    ("Facturas", 0),
                    ("Notas de Crédito", 0),
                    ("Notas de Débito", 0),
                    ("Comprobantes de Crédito Fiscal", 1),
                    ("Documentos Procesados", 1),
                ],
            )

            libro = load_workbook(carpeta / "compras.xlsx")
            self.assertEqual(libro.sheetnames, ["Resumen de DTEs", "Percepción", "Errores"])

            fila_resumen = list(libro["Resumen de DTEs"].values)[1]
            self.assertEqual(fila_resumen[0], "21/06/2026")
            self.assertEqual(fila_resumen[2], "03")
            self.assertEqual(fila_resumen[6], 11.00)
            self.assertEqual(fila_resumen[9], 15.00)
            self.assertEqual(fila_resumen[13], 3.00)

            fila_percepcion = list(libro["Percepción"].values)[1]
            self.assertEqual(fila_percepcion[3], "SELLO-001")
            self.assertEqual(fila_percepcion[6], 0.30)

            libro_errores = load_workbook(carpeta / "compras_errores.xlsx")
            self.assertEqual(list(libro_errores["Errores"].values)[1], ("nota.txt", "No es un archivo JSON"))
            self.assertEqual(len(progresos), 2)
            self.assertTrue(all(total == 2 for _, total, _ in progresos))

    def test_excluye_dte_invalidado_por_hacienda(self):
        with tempfile.TemporaryDirectory() as directorio:
            carpeta = Path(directorio)
            dte = {
                "identificacion": {
                    "fecEmi": "2026-05-26",
                    "tipoDte": "03",
                    "codigoGeneracion": "DTE-INVALIDADO",
                },
                "emisor": {"nit": "0614-000000-000-0", "nrc": "1", "nombre": "Proveedor"},
                "resumen": {"montoTotalOperacion": 10.00},
            }
            (carpeta / "invalidado.json").write_text(json.dumps(dte), encoding="utf-8")

            resumen = generar_anexo_compras(
                carpeta,
                carpeta,
                "compras",
                consultar_estado=lambda codigo, fecha: {
                    "estadoDoc": "Invalidado",
                    "descripcionEstado": "Documento Invalidado por el Emisor.",
                },
            )

            self.assertEqual(resumen[-1], ("Documentos Procesados", 0))
            libro = load_workbook(carpeta / "compras.xlsx")
            self.assertEqual(libro["Resumen de DTEs"].max_row, 1)

            self.assertFalse((carpeta / "compras_errores.xlsx").exists())
            libro_invalidos = load_workbook(carpeta / "InvalidosDTE.xlsx")
            self.assertEqual(
                list(libro_invalidos["DTE Invalidos"].values)[1],
                (
                    "invalidado.json",
                    "DTE-INVALIDADO",
                    "26/05/2026",
                    "Invalidado",
                    "Documento Invalidado por el Emisor.",
                ),
            )

    def test_procesa_dte_si_hacienda_alcanza_el_limite_de_consultas(self):
        with tempfile.TemporaryDirectory() as directorio:
            carpeta = Path(directorio)
            dte = {
                "identificacion": {
                    "fecEmi": "2026-09-07",
                    "tipoDte": "03",
                    "codigoGeneracion": "GEN-SIN-VERIFICAR",
                },
                "emisor": {
                    "nit": "0614-000000-000-0",
                    "nrc": "12345",
                    "nombre": "Proveedor de Prueba",
                },
                "resumen": {"montoTotalOperacion": 25.00},
                "selloRecibido": "SELLO-LOCAL",
            }
            (carpeta / "sin_verificar.json").write_text(
                json.dumps(dte),
                encoding="utf-8",
            )

            def consulta_limitada(codigo, fecha):
                raise LimiteConsultasHaciendaError("HTTP 429")

            resumen = generar_anexo_compras(
                carpeta,
                carpeta,
                "compras",
                consultar_estado=consulta_limitada,
            )

            self.assertEqual(resumen[-1], ("Documentos Procesados", 1))
            libro = load_workbook(carpeta / "compras.xlsx")
            self.assertEqual(libro["Resumen de DTEs"].max_row, 2)

            libro_errores = load_workbook(carpeta / "compras_errores.xlsx")
            aviso = list(libro_errores["Errores"].values)[1]
            self.assertEqual(aviso[0], "sin_verificar.json")
            self.assertIn("Procesado sin verificación", aviso[1])
