import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from reports.anexos_retencion import generar_anexo_retencion


class AnexosRetencionTests(unittest.TestCase):
    def test_genera_anexo_de_retencion_y_archivo_de_errores(self):
        with tempfile.TemporaryDirectory() as directorio:
            carpeta = Path(directorio)
            dte = {
                "identificacion": {
                    "fecEmi": "2026-06-21",
                    "tipoDte": "07",
                    "codigoGeneracion": "RET-001",
                },
                "emisor": {"nit": "0614-000000-000-0"},
                "resumen": {
                    "totalSujetoRetencion": 100.00,
                    "totalIVAretenido": 13.00,
                },
                "responseMH": {"selloRecibido": "SELLO-RET"},
            }
            (carpeta / "retencion.json").write_text(json.dumps(dte), encoding="utf-8")
            (carpeta / "invalido.json").write_text("{invalido", encoding="utf-8")

            resumen = generar_anexo_retencion(
                carpeta,
                carpeta,
                "retenciones",
                consultar_estado=lambda codigo, fecha: {
                    "estadoDoc": "Transmitido Satisfactoriamente",
                    "selloVal": "SELLO-HACIENDA",
                },
            )

            self.assertEqual(
                resumen,
                [("Comprobantes de Retención", 1), ("Documentos Procesados", 1)],
            )

            libro = load_workbook(carpeta / "retenciones.xlsx")
            self.assertEqual(libro.sheetnames, ["Comprobantes de Retención", "Errores"])
            fila = list(libro["Comprobantes de Retención"].values)[1]
            self.assertEqual(fila[1], "21/06/2026")
            self.assertEqual(fila[3], "SELLO-HACIENDA")
            self.assertEqual(fila[5], 100.00)
            self.assertEqual(fila[6], 13.00)

            libro_errores = load_workbook(carpeta / "retenciones_errores.xlsx")
            self.assertEqual(list(libro_errores["Errores"].values)[1], ("invalido.json", "JSON inválido"))

    def test_excluye_retencion_no_transmitida_satisfactoriamente(self):
        with tempfile.TemporaryDirectory() as directorio:
            carpeta = Path(directorio)
            dte = {
                "identificacion": {
                    "fecEmi": "2026-05-27",
                    "tipoDte": "07",
                    "codigoGeneracion": "RET-INVALIDADA",
                },
                "emisor": {"nit": "0614-000000-000-0"},
                "resumen": {},
            }
            (carpeta / "retencion_invalidada.json").write_text(json.dumps(dte), encoding="utf-8")

            resumen = generar_anexo_retencion(
                carpeta,
                carpeta,
                "retenciones",
                consultar_estado=lambda codigo, fecha: {
                    "estadoDoc": "Invalidado",
                    "descripcionEstado": "Documento Invalidado por el Emisor.",
                },
            )

            self.assertEqual(
                resumen,
                [("Comprobantes de Retención", 0), ("Documentos Procesados", 0)],
            )
            libro = load_workbook(carpeta / "retenciones.xlsx")
            self.assertEqual(libro["Comprobantes de Retención"].max_row, 1)

            libro_invalidos = load_workbook(carpeta / "InvalidosDTE.xlsx")
            self.assertEqual(
                list(libro_invalidos["DTE Invalidos"].values)[1],
                (
                    "retencion_invalidada.json",
                    "RET-INVALIDADA",
                    "27/05/2026",
                    "Invalidado",
                    "Documento Invalidado por el Emisor.",
                ),
            )
