import copy
import json
import tempfile
import unittest
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import load_workbook

import libroIVA
from dte.types import (
    TIPO_LIBRO_COMPRAS,
    TIPO_LIBRO_VENTAS_CONSUMIDOR_FINAL,
    TIPO_LIBRO_VENTAS_CONTRIBUYENTES,
)
from tests.dte_fixtures import crear_dte


class LibrosIVATests(unittest.TestCase):
    def setUp(self):
        self.recursos = ExitStack()
        self.addCleanup(self.recursos.close)
        self.directorio = Path(self.recursos.enter_context(tempfile.TemporaryDirectory()))
        self.entrada = self.directorio / "entrada"
        self.entrada.mkdir()
        self.salida = self.directorio / "libro.xlsx"
        # Sustituir solamente los controles de la GUI; generar el Excel real.
        controles = {
            "nombre_contri_var": "Contribuyente Prueba",
            "nrc_entry": "123",
            "nit_entry": "06140000000000",
            "dui_entry": "",
            "month_combobox": "Octubre",
            "year_combobox": "2026",
            "nombre_firma_var": "Contador Prueba",
        }
        for nombre, valor in controles.items():
            self.recursos.enter_context(patch.object(
                libroIVA, nombre, Mock(get=Mock(return_value=valor)), create=True,
            ))
        self.recursos.enter_context(patch.object(libroIVA.messagebox, "showinfo"))

    def escribir(self, nombre, dte):
        (self.entrada / nombre).write_text(json.dumps(dte), encoding="utf-8")

    def generar(self, tipo):
        libroIVA.generar_libro_excel(self.entrada, self.salida, tipo)
        libro = load_workbook(self.salida)
        self.addCleanup(libro.close)
        return libro

    def comprobar_totales(self, hoja, cantidad_filas, ultima_columna):
        fila_total = 14 + cantidad_filas
        self.assertEqual(hoja.cell(fila_total, 5).value, "Totales:")
        for columna in range(6, ultima_columna + 1):
            suma_filas = sum(
                float(hoja.cell(fila, columna).value or 0)
                for fila in range(13, 13 + cantidad_filas)
            )
            self.assertAlmostEqual(hoja.cell(fila_total, columna).value, suma_filas)
        return fila_total

    def test_compra_del_ejemplo_suma_no_sujetas_y_exentas_en_total(self):
        self.escribir("compra.json", crear_dte())
        libro = self.generar(TIPO_LIBRO_COMPRAS)
        hoja = libro.active
        self.assertEqual(hoja["F13"].value, 4)
        self.assertEqual(hoja["H13"].value, 3)
        self.assertAlmostEqual(hoja["K13"].value, 0.39)
        self.assertAlmostEqual(hoja["L13"].value, 7.39)
        fila_total = self.comprobar_totales(hoja, 1, 14)
        self.assertEqual(hoja.cell(fila_total, 6).value, 4)
        self.assertAlmostEqual(hoja.cell(fila_total, 12).value, 7.39)
        self.assertNotIn("Avisos", libro.sheetnames)

    def test_compras_y_contribuyentes_restan_notas_credito_y_omiten_copias(self):
        # Dos CCF distintos con igual fecha e importes deben conservarse.
        self.escribir("01-ccf.json", crear_dte("03", "CCF-1", factor=2))
        self.escribir("02-ccf.json", crear_dte("03", "CCF-2", factor=2))
        self.escribir("03-credito.json", crear_dte("05", "NC-1"))
        self.escribir("04-debito.json", crear_dte("06", "ND-1"))
        copia = crear_dte("05", " nc-1 ")
        self.escribir("05-copia.json", {"dte": copia})
        for tipo in (TIPO_LIBRO_COMPRAS, TIPO_LIBRO_VENTAS_CONTRIBUYENTES):
            with self.subTest(tipo=tipo):
                libro = self.generar(tipo)
                hoja = libro.active
                fila_total = self.comprobar_totales(hoja, 4, 14)
                self.assertEqual(hoja["F15"].value, -4)
                self.assertEqual(hoja.cell(fila_total, 6).value, 16)
                self.assertAlmostEqual(hoja.cell(fila_total, 12).value, 29.56)
                avisos = list(libro["Avisos"].values)
                self.assertEqual(len(avisos), 2)
                self.assertEqual(avisos[1][0], "05-copia.json")
                self.assertIn("03-credito.json", avisos[1][1])
                if tipo == TIPO_LIBRO_VENTAS_CONTRIBUYENTES:
                    # Resumen de operaciones y total de la columna deben coincidir.
                    self.assertEqual(hoja.cell(fila_total + 5, 10).value, 16)
                libro.close()

    def test_consumidor_final_no_duplica_importes_agrupados_y_concilia_exportacion(self):
        factura = crear_dte("01", "FAC-1")
        self.escribir("01-factura.json", factura)
        segunda = crear_dte("01", "FAC-2")
        segunda["identificacion"]["horEmi"] = "09:00:00"
        self.escribir("02-factura.json", segunda)
        self.escribir("03-copia.json", {"dte": copy.deepcopy(factura)})
        self.escribir("04-exportacion.json", crear_dte("11", "EXP-1"))
        libro = self.generar(TIPO_LIBRO_VENTAS_CONSUMIDOR_FINAL)
        hoja = libro.active
        fila_total = self.comprobar_totales(hoja, 2, 11)
        self.assertEqual(hoja["C13"].value, "FAC-1")
        self.assertEqual(hoja["D13"].value, "FAC-2")
        self.assertEqual(hoja["F13"].value, 4)
        self.assertEqual(hoja["G13"].value, 6)
        self.assertEqual(hoja["H14"].value, 3)
        self.assertEqual(hoja.cell(fila_total, 6).value, 4)
        self.assertAlmostEqual(hoja.cell(fila_total, 11).value, 22.17)
        self.assertEqual(hoja.cell(fila_total + 6, 7).value, 4)
        self.assertEqual(libro["Avisos"].max_row, 2)

    def test_consumidor_final_acepta_receptor_null_y_conserva_importes(self):
        factura = crear_dte("01", "FAC-SIN-RECEPTOR")
        factura["receptor"] = None
        factura["resumen"].update({
            "totalNoSuj": 0, "totalNoGravado": 0,
            "totalGravada": 113, "totalExenta": 2,
            "montoTotalOperacion": 115, "ivaRete1": 1,
        })
        for anidado in (False, True):
            with self.subTest(anidado=anidado):
                self.escribir("factura.json", {"dte": factura} if anidado else factura)
                libro = self.generar(TIPO_LIBRO_VENTAS_CONSUMIDOR_FINAL)
                hoja = libro.active
                self.assertEqual(hoja["C13"].value, "FAC-SIN-RECEPTOR")
                self.assertEqual(hoja["D13"].value, "FAC-SIN-RECEPTOR")
                # openpyxl lee una celda guardada con texto vacío como None.
                self.assertIsNone(hoja["E13"].value)
                self.assertEqual(
                    [hoja.cell(13, columna).value for columna in range(6, 12)],
                    [2, 113, 0, 0, 1, 115],
                )
                fila_total = self.comprobar_totales(hoja, 1, 11)
                self.assertEqual(
                    [hoja.cell(fila_total + 6, columna).value for columna in range(7, 12)],
                    [2, 100, 13, 0, 115],
                )
                libro.close()

    def test_descuentos_percepcion_retencion_y_tributos_concilian_por_columna(self):
        dte = crear_dte()
        dte["resumen"].update({
            "descuExenta": 0.5, "descuNoSuj": 0.25, "descuGravada": 1,
            "ivaPerci1": 0.03, "ivaRete1": 0.02,
            "tributos": [
                {"codigo": "20", "valor": 0.26},
                {"codigo": "D1", "valor": 0.10},
                {"codigo": "C8", "valor": 0.05},
            ],
        })
        self.escribir("compra.json", dte)
        for tipo in (TIPO_LIBRO_COMPRAS, TIPO_LIBRO_VENTAS_CONTRIBUYENTES):
            with self.subTest(tipo=tipo):
                libro = self.generar(tipo)
                fila_total = self.comprobar_totales(libro.active, 1, 14)
                self.assertEqual(libro.active.cell(fila_total, 6).value, 3.25)
                self.assertAlmostEqual(libro.active.cell(fila_total, 13).value, 0.03)
                if tipo == TIPO_LIBRO_COMPRAS:
                    self.assertAlmostEqual(libro.active.cell(fila_total, 10).value, 0.15)
                else:
                    self.assertAlmostEqual(libro.active.cell(fila_total, 14).value, 0.02)
                libro.close()

    def test_libros_vacios_generan_totales_cero(self):
        for tipo, ultima_columna in (
            (TIPO_LIBRO_COMPRAS, 14),
            (TIPO_LIBRO_VENTAS_CONTRIBUYENTES, 14),
            (TIPO_LIBRO_VENTAS_CONSUMIDOR_FINAL, 11),
        ):
            with self.subTest(tipo=tipo):
                libro = self.generar(tipo)
                self.comprobar_totales(libro.active, 0, ultima_columna)
                libro.close()
