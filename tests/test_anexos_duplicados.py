import copy
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock

from openpyxl import load_workbook

from reports.anexos_anulados import generar_anexo_anulados
from reports.anexos_compras import generar_anexo_compras
from reports.anexos_consumidor_final import generar_anexo_consumidor_final
from reports.anexos_contribuyentes import generar_anexo_contribuyentes
from reports.anexos_retencion import generar_anexo_retencion
from reports.anexos_sujeto_excluido import generar_anexo_sujeto_excluido
from tests.dte_fixtures import crear_dte


class AnexosDuplicadosTests(unittest.TestCase):
    def test_los_seis_anexos_omiten_copias_sin_alterar_filas_importes_o_contadores(self):
        casos = (
            (generar_anexo_compras, "03", True),
            (generar_anexo_contribuyentes, "03", False),
            (generar_anexo_consumidor_final, "01", False),
            (generar_anexo_retencion, "07", True),
            (generar_anexo_sujeto_excluido, "14", False),
            (generar_anexo_anulados, "03", False),
        )
        for generar, tipo, consulta_hacienda in casos:
            with self.subTest(anexo=generar.__name__), tempfile.TemporaryDirectory() as directorio:
                raiz = Path(directorio)
                entrada = raiz / "entrada"
                entrada.mkdir()
                # Igual fecha e importes, pero distinto código: ambos se conservan.
                primero = crear_dte(tipo, "GEN-1")
                segundo = crear_dte(tipo, "GEN-2")
                primero["resumen"]["ivaPerci1"] = 0.03
                segundo["resumen"]["ivaPerci1"] = 0.03
                for nombre, dte in (("01-original.json", primero), ("02-distinto.json", segundo)):
                    (entrada / nombre).write_text(json.dumps(dte), encoding="utf-8")

                consulta = Mock(return_value={"estadoDoc": "Transmitido Satisfactoriamente"})
                argumentos = {"consultar_estado": consulta} if consulta_hacienda else {}
                resumen_original = generar(entrada, raiz, "original", **argumentos)
                if consulta_hacienda:
                    self.assertEqual(consulta.call_count, 2)
                consulta.reset_mock()

                copia = copy.deepcopy(primero)
                copia["identificacion"]["codigoGeneracion"] = " gen-1 "
                (entrada / "03-copia.json").write_text(json.dumps({"dte": copia}), encoding="utf-8")
                resumen_copia = generar(entrada, raiz, "con_copia", **argumentos)

                self.assertEqual(resumen_copia, resumen_original)
                self.assertEqual(dict(resumen_copia)["Documentos Procesados"], 2)
                if consulta_hacienda:
                    self.assertEqual(consulta.call_count, 2, "La copia no debe consultar Hacienda")

                original = load_workbook(raiz / "original.xlsx")
                con_copia = load_workbook(raiz / "con_copia.xlsx")
                errores = load_workbook(raiz / "con_copia_errores.xlsx")
                try:
                    for hoja in original.sheetnames:
                        if hoja != "Errores":
                            self.assertEqual(list(original[hoja].values), list(con_copia[hoja].values))
                    avisos = list(con_copia["Errores"].values)
                    self.assertEqual(len(avisos), 2)
                    self.assertEqual(avisos[1][0], "03-copia.json")
                    self.assertIn("duplicado omitido", avisos[1][1])
                    self.assertIn("01-original.json", avisos[1][1])
                    self.assertEqual(list(errores["Errores"].values), avisos)
                finally:
                    original.close()
                    con_copia.close()
                    errores.close()
